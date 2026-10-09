"""Run the Check over a list of companies and measure it.

Inputs:
  data/test_companies.json        public test set: name, url, about?, expected (head_on | supplier | open)
  a PitchBook export (.csv/.xlsx)  private; results stay in data/private/

Measures: agreement between the Check's position and the expected label, how often cited evidence ids
exist, and time per check. The first call runs alone to warm the prompt cache; the rest run in parallel.

Usage:
  .venv/bin/python -m gt.evaluate data/test_companies.json
  .venv/bin/python -m gt.evaluate data/private/pitchbook.csv --limit 40
"""
import argparse
import csv
import json
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from gt import calls
from gt.check import MODEL, check, cited_ids_ok, slugify
from gt.env import DATA

FALLBACK = "claude-opus-5-5"
EXPECTED = {"head_on": {"head_on_with_entrenched", "head_on_with_building"},
            "supplier": {"complement"}, "open": {"white_space"}}


def from_table(path):
    """A PitchBook export (.csv or .xlsx). Exports start with banner rows, so find the header row first,
    then the name, website and description columns by their headers."""
    path = Path(path)
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook
        sheet = load_workbook(path, read_only=True, data_only=True).active
        table = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            table = list(csv.reader(f))
    start = next(i for i, row in enumerate(table) if any("compan" in c.lower() for c in row) and len([c for c in row if c]) >= 3)
    headers = [h.strip() for h in table[start]]
    rows = [dict(zip(headers, r)) for r in table[start + 1:] if any(r)]

    def col(*words):
        return next((h for w in words for h in headers if w in h.lower()), None)

    name, site, about = col("companies", "company name", "company"), col("website", "url"), col("description")
    return [{"name": r[name].strip(), "url": (r.get(site) or "").strip(), "about": (r.get(about) or "").strip()}
            for r in rows if r.get(name, "").strip()]


def file_names(companies):
    """One saved file per company. Names repeat (the PitchBook export has two different companies with one name),
    so a repeated name gets its website added; the first keeps the plain name, so earlier runs still resume."""
    seen, names = set(), []
    for c in companies:
        name = slugify(c["name"])
        if name in seen:
            name = slugify(f"{c['name']} {c.get('url', '').split('//')[-1]}")
        seen.add(name)
        names.append(name)
    return names


def run_one(company, model=None):
    url = company.get("url", "")
    if url and not url.startswith("http"):
        url = "https://" + url
    started = time.time()
    note = None
    try:
        verdict, usage = check(company["name"], url, company.get("about", ""), model=model)
    except Exception as e:  # keep going; a failed check is a result too
        # Haiku has no server-side refusal fallback, and its safety filter sometimes stops on harmless companies
        # (crop-spraying drones). A refusal, or an answer cut off mid-JSON, is redone on Opus and recorded.
        if (model or MODEL) == FALLBACK or not ("refusal" in str(e) or "Invalid JSON" in str(e)):
            return {**company, "error": str(e)[:200]}
        note = f"{model or MODEL} gave no verdict ({str(e)[:80]}); redone on {FALLBACK}."
        try:
            verdict, usage = check(company["name"], url, company.get("about", ""), model=FALLBACK)
        except Exception as e2:
            return {**company, "error": f"{note} {e2}"[:200]}
    result = {**company, "verdict": verdict.model_dump(), "unknown_citations": cited_ids_ok(verdict),
              **calls.money(verdict), "usage": usage, "seconds": round(time.time() - started, 1)}
    return {**result, "note": note} if note else result


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--workers", type=int, default=4)
    ap.add_argument("--model", default=None, help="compare another Claude model (default: gt.check.MODEL)")
    ap.add_argument("--out", default=None, help="save to this directory instead of data/checks or data/private/checks")
    args = ap.parse_args()
    src = Path(args.source)
    companies = from_table(src) if src.suffix in (".csv", ".xlsx") else json.loads(src.read_text())
    companies = companies[: args.limit] if args.limit else companies
    private = "private" in src.parts
    out_dir = Path(args.out) if args.out else DATA / ("private/checks" if private else "checks")
    if args.model and args.model != MODEL:
        out_dir = out_dir / f"compare-{args.model}"
    out_dir.mkdir(parents=True, exist_ok=True)

    def run(company, name):
        """Each verdict is saved as soon as it arrives, and a rerun skips what's already saved."""
        path = out_dir / f"{name}.json"
        if path.exists():
            return {**company, **json.loads(path.read_text())}
        r = run_one(company, args.model)
        if "verdict" in r:
            path.write_text(json.dumps(r, indent=1) + "\n")
        return r

    names = file_names(companies)
    results = [run(companies[0], names[0])]  # warm the prompt cache first
    with ThreadPoolExecutor(args.workers) as pool:
        results += list(pool.map(run, companies[1:], names[1:]))

    done = [r for r in results if "verdict" in r]
    labeled = [r for r in done if r.get("expected") in EXPECTED]
    agree = [r for r in labeled if r["verdict"]["position"] in EXPECTED[r["expected"]]]
    cited = sum(len(o["evidence_ids"]) for r in done for o in r["verdict"]["overlaps"])
    bad = sum(len(r["unknown_citations"]) for r in done)
    seconds = sorted(r["seconds"] for r in done if r.get("seconds"))  # a hand rerun has no timing
    report = {
        "checked": len(done), "failed": len(results) - len(done),
        "agreement": f"{len(agree)}/{len(labeled)}" if labeled else "no labels",
        "positions": Counter(r["verdict"]["position"] for r in done),
        "calls": Counter(r["verdict"]["call"] for r in done),
        "citations": {"total": cited, "unknown": bad},
        "median_seconds": seconds[len(seconds) // 2] if seconds else None,
        "disagreements": [{"name": r["name"], "expected": r["expected"], "got": r["verdict"]["position"],
                           "reasoning": r["verdict"]["reasoning"]} for r in labeled if r not in agree],
        "errors": [{"name": r["name"], "error": r["error"]} for r in results if "error" in r],
    }
    (out_dir / "_report.json").write_text(json.dumps(report, indent=1, default=dict) + "\n")
    print(json.dumps(report, indent=1, default=dict))


if __name__ == "__main__":
    main()
